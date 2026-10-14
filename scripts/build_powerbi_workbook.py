"""Build the small Power BI Service fallback workbook from a curated snapshot.

Artifact Tool is preferred for spreadsheet authoring, but unavailable in this
environment; openpyxl is the explicit fallback. Only masked curated rows enter.
"""

import argparse
import json
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

SHEETS = [
    "fact_fraud_daily",
    "fact_fraud_decisions",
    "fact_alert_performance",
    "fact_customer_timeline",
    "fact_pipeline_health",
    "dim_date",
    "dim_transaction_type",
    "dim_rule",
]


def typed(value: object, column: str) -> object:
    if isinstance(value, str) and column in {"event_date"}:
        return date.fromisoformat(value)
    if isinstance(value, str) and column.endswith(("_at", "_time")):
        parsed = datetime.fromisoformat(value)
        return parsed.astimezone(UTC).replace(tzinfo=None) if parsed.tzinfo else parsed
    return value


def build(snapshot: Path) -> Path:
    manifest = json.loads((snapshot / "manifest.json").read_text())
    workbook = Workbook()
    info = workbook.active
    info.title = "Snapshot Info"
    info.append(["Field", "Value"])
    for key in ["snapshot_id", "exported_at", "run_id", "semantic_note"]:
        info.append([key, manifest[key]])
    info.append(["Data scope", "100k Gold transactions; one selected rules run"])
    info.append(["Privacy", "customer_key is HMAC masked; raw IDs excluded"])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 74

    for name in SHEETS:
        rows = json.loads((snapshot / f"{name}.json").read_text())
        columns = list(rows[0])
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append([typed(row.get(column), column) for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.alignment = Alignment(horizontal="center")
        for index, column in enumerate(columns, start=1):
            letter = sheet.cell(1, index).column_letter
            sheet.column_dimensions[letter].width = min(34, max(14, len(column) + 2))
            if column == "event_date":
                for cell in sheet[letter][1:]:
                    cell.number_format = "yyyy-mm-dd"
            elif column.endswith(("_at", "_time")):
                for cell in sheet[letter][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
        table = Table(displayName="T_" + name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    output = snapshot / "powerbi_fraud_snapshot.xlsx"
    workbook.save(output)
    # Verify every imported table and masking before offering it to Power BI.
    reopened = load_workbook(output, read_only=False, data_only=True)
    for name in SHEETS:
        assert reopened[name].max_row - 1 == manifest["tables"][name]["rows"]
    assert "customer_id" not in [cell.value for cell in reopened["fact_fraud_decisions"][1]]
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("snapshot", type=Path)
    args = parser.parse_args()
    print(build(args.snapshot))


if __name__ == "__main__":
    main()
